"""
"홍산디_2차년도_3회차_전기전자제품_메타데이터.xlsx"를 IndustryDataRequest 모양의 dict 리스트로 파싱한다.

엑셀 구조:
- row 1: 헤더 (ID, image name, company, model, category, type, usage, weight_kg, size_mm, price_krw, release_date, store_url)
- 그룹 헤더 행: A열만 채워짐 (예: "커피/에스프레소 머신") - 그 아래 데이터의 소분류를 나타냄
- 그 외 행: 실제 데이터
"""
import argparse
import json
import sys

import openpyxl

GROUP_TO_CATEGORY = {
    "커피/에스프레소 머신": "COFFEE_MACHINE",
    "커피메이커": "COFFEE_MAKER",
    "원두 전동 그라인더": "COFFEE_GRINDER",
    "압력밥솥": "PRESSURE_RICE_COOKER",
    "멀티쿠커": "MULTI_COOKER",
    "에어프라이어": "AIR_FRYER",
    "믹서기": "MIXER",
}


def to_str(value):
    if value is None:
        return None
    if isinstance(value, float):
        if value == int(value):
            return str(int(value))
        return str(value)
    return str(value).strip() or None


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["최종"]

    rows = []
    current_category = None

    for r in range(2, ws.max_row + 1):
        cell_a = ws.cell(row=r, column=1).value
        cell_b = ws.cell(row=r, column=2).value

        # 그룹 헤더 행: A열만 있고 B열이 비어있음
        if cell_a is not None and cell_b is None:
            group_name = to_str(cell_a)
            if group_name not in GROUP_TO_CATEGORY:
                raise ValueError(f"알 수 없는 그룹명: {group_name!r} (row {r})")
            current_category = GROUP_TO_CATEGORY[group_name]
            continue

        if cell_a is None:
            continue

        if current_category is None:
            raise ValueError(f"카테고리 그룹 헤더보다 먼저 데이터가 나옴 (row {r})")

        code = to_str(cell_a)
        image_name = to_str(ws.cell(row=r, column=2).value)
        company = to_str(ws.cell(row=r, column=3).value)
        model = to_str(ws.cell(row=r, column=4).value)
        category_path = to_str(ws.cell(row=r, column=5).value)
        product_type = to_str(ws.cell(row=r, column=6).value)
        usage = to_str(ws.cell(row=r, column=7).value)
        weight_kg = to_str(ws.cell(row=r, column=8).value)
        size_mm = to_str(ws.cell(row=r, column=9).value)
        price_krw = to_str(ws.cell(row=r, column=10).value)
        release_date = to_str(ws.cell(row=r, column=11).value)
        store_url = to_str(ws.cell(row=r, column=12).value)

        product_name = " ".join(p for p in [company, model] if p) or None

        rows.append({
            "code": code,
            "imageName": image_name,
            "productName": product_name,
            "companyName": company,
            "modelName": model,
            "price": price_krw,
            "material": None,
            "size": size_mm,
            "weight": weight_kg,
            "referenceUrl": store_url,
            "registeredAt": release_date,
            "productPath": category_path,
            "productTypeName": product_type,
            "noiseCancelling": None,
            "codec": None,
            "extraFeatures": None,
            "controlType": None,
            "waterproof": None,
            "maxPlayTime": None,
            "chargeTime": None,
            "usage": usage,
            "shoppingUrl": None,
            "soundOutput": None,
            "connectivity": None,
            "originalDetailImagePath": None,
            "originalFrontImagePath": None,
            "originalSideImagePath": None,
            "originalSide2ImagePath": None,
            "originalSide3ImagePath": None,
            "industryDataCategory": current_category,
        })

    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("excel_path")
    parser.add_argument("--out", default="industry_data.json")
    parser.add_argument("--limit", type=int, default=None)
    args = parser.parse_args()

    rows = parse(args.excel_path)
    if args.limit:
        rows = rows[:args.limit]

    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)

    print(f"파싱 완료: {len(rows)}건 -> {args.out}", file=sys.stderr)


if __name__ == "__main__":
    main()
